from __future__ import annotations

from pathlib import Path
import mimetypes

DOCUMENT_EXTENSIONS = {".pdf", ".ppt", ".pptx"}
DOCUMENT_MIME_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.ms-powerpoint": ".ppt",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
}


def is_supported_document(value: str | Path) -> bool:
    suffix = Path(str(value)).suffix.lower()
    return suffix in DOCUMENT_EXTENSIONS


def get_document_kind(value: str | Path) -> str:
    suffix = Path(str(value)).suffix.lower()
    return {".pdf": "pdf", ".ppt": "ppt", ".pptx": "pptx"}.get(suffix, "document")


def document_mime_type(path: str | Path) -> str:
    suffix = Path(str(path)).suffix.lower()
    explicit = {
        ".pdf": "application/pdf",
        ".ppt": "application/vnd.ms-powerpoint",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    return explicit.get(suffix) or mimetypes.guess_type(str(path))[0] or "application/octet-stream"


def inspect_document(path: str | Path) -> dict:
    """Return lightweight document metadata without converting the document.

    PDF page counts and PPTX slide counts are read when the optional parsers are
    installed. The Canva import path itself does not depend on these parsers.
    """
    file_path = Path(path)
    result = {
        "name": file_path.name,
        "kind": get_document_kind(file_path),
        "mime_type": document_mime_type(file_path),
        "size": file_path.stat().st_size if file_path.exists() else 0,
        "page_count": None,
        "slide_count": None,
        "text_preview": "",
    }

    if not file_path.exists():
        return result

    if file_path.suffix.lower() == ".pdf":
        try:
            import fitz  # PyMuPDF
            with fitz.open(file_path) as doc:
                result["page_count"] = len(doc)
                chunks = []
                for page in doc[: min(3, len(doc))]:
                    text = page.get_text("text").strip()
                    if text:
                        chunks.append(text)
                result["text_preview"] = "\n\n".join(chunks)[:4000]
        except Exception:
            pass
    elif file_path.suffix.lower() == ".pptx":
        try:
            from pptx import Presentation
            presentation = Presentation(str(file_path))
            result["slide_count"] = len(presentation.slides)
            chunks = []
            for slide in presentation.slides[:3]:
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
                if texts:
                    chunks.append("\n".join(texts))
            result["text_preview"] = "\n\n".join(chunks)[:4000]
        except Exception:
            pass

    return result